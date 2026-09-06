from datetime import UTC, datetime
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from sqlalchemy import asc, desc, or_
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from app.database import get_db
from app.models.category import Category
from app.schemas.category import CategoryCreate, CategoryResponse, CategoryUpdate
from app.schemas.common import ok
from app.services.auth import get_current_user

router = APIRouter(prefix="/categories", tags=["categories"])


@router.get("")
def list_categories(
    db: Session = Depends(get_db),
    _user=Depends(get_current_user),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    search: str | None = Query(default=None),
    sort_by: str | None = Query(default="name"),
    sort_dir: str | None = Query(default="asc"),
):
    query = db.query(Category).filter(Category.deleted_at.is_(None))
    if search:
        like = f"%{search}%"
        query = query.filter(or_(Category.name.ilike(like), Category.slug.ilike(like)))
    sort_col = getattr(Category, sort_by or "name", Category.name)
    query = query.order_by(sort_col.desc() if sort_dir == "desc" else sort_col.asc())
    total = query.count()
    items = query.offset((page - 1) * page_size).limit(page_size).all()
    return ok([CategoryResponse.model_validate(i).model_dump(mode="json") for i in items], meta={"page": page, "page_size": page_size, "total": total})


@router.post("")
def create_category(payload: CategoryCreate, db: Session = Depends(get_db), _user=Depends(get_current_user)):
    if payload.slug:
        exists = db.query(Category).filter(Category.slug == payload.slug, Category.deleted_at.is_(None)).first()
        if exists:
            raise HTTPException(status_code=400, detail="Slug already exists")
    if payload.parent_id is not None:
        parent = db.get(Category, payload.parent_id)
        if parent is None or parent.deleted_at is not None:
            raise HTTPException(status_code=400, detail="Parent category not found")
    category = Category(**payload.model_dump())
    db.add(category)
    db.commit()
    db.refresh(category)
    return ok(CategoryResponse.model_validate(category).model_dump(mode="json"), meta={"id": category.id})


@router.post("/import")
def import_categories(file: UploadFile = File(...), db: Session = Depends(get_db), _user=Depends(get_current_user)):
    """Import categories from an Excel file.

    Expected columns (first row = header):
      name, slug, parent, active
    - `parent` is matched by category name.
    - Rows with an existing slug are skipped as duplicates.
    - `active` accepts 1/0, "yes"/"no", "true"/"false", or leave blank (defaults active).
    """
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Only .xlsx files are supported")
    try:
        wb = load_workbook(BytesIO(file.file.read()), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read the Excel file")
    ws = wb.active
    wb.close()
    if ws is None:
        raise HTTPException(status_code=400, detail="Excel workbook has no active sheet")
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="Excel file is empty")

    header = [str(c or "").strip().lower() if c is not None else "" for c in rows[0]]
    def col_idx(*names: str) -> int | None:
        for i, h in enumerate(header):
            if h in names:
                return i
        return None

    name_idx = col_idx("name", "نام")
    if name_idx is None:
        raise HTTPException(status_code=400, detail='Missing required column "name"')

    slug_idx = col_idx("slug", "اسلاگ")
    parent_idx = col_idx("parent", "والد", "parent_name", "parent name")
    active_idx = col_idx("active", "فعال")

    created, skipped, failed = 0, 0, 0
    errors: list[dict] = []
    for row_num, row in enumerate(rows[1:], start=2):
        name = str(row[name_idx] or "").strip() if name_idx is not None and row[name_idx] is not None else ""
        if not name:
            failed += 1
            errors.append({"row": row_num, "reason": "name is required"})
            continue
        slug = str(row[slug_idx] or "").strip() if slug_idx is not None and row[slug_idx] is not None else None
        if slug:
            # Check any row, including soft-deleted ones: the DB has a UNIQUE
            # constraint on slug, so a soft-deleted row with the same slug
            # would still block the insert.
            existing = db.query(Category).filter(Category.slug == slug).first()
            if existing:
                skipped += 1
                continue
        parent = None
        if parent_idx is not None and row[parent_idx]:
            parent_name = str(row[parent_idx]).strip()
            parent = db.query(Category).filter(
                Category.name == parent_name, Category.deleted_at.is_(None)
            ).first()
            if parent is None:
                failed += 1
                errors.append({"row": row_num, "reason": f'parent "{parent_name}" not found'})
                continue
        active = True
        if active_idx is not None and row[active_idx] is not None:
            val = str(row[active_idx]).strip().lower()
            active = val in ("1", "true", "yes", "بله", "فعال")
        category = Category(name=name, slug=slug or None, parent_id=parent.id if parent else None, active=active)
        db.add(category)
        # Flush so categories created earlier in this file are visible to the
        # parent lookup of subsequent rows (session has autoflush disabled).
        db.flush()
        created += 1
    db.commit()
    return ok(
        {"created": created, "skipped": skipped, "failed": failed, "errors": errors},
        meta={"total_rows": len(rows) - 1},
    )


@router.get("/{category_id}")
def get_category(category_id: int, db: Session = Depends(get_db), _user=Depends(get_current_user)):
    category = db.get(Category, category_id)
    if category is None or category.deleted_at is not None:
        raise HTTPException(status_code=404, detail="Category not found")
    return ok(CategoryResponse.model_validate(category).model_dump(mode="json"))


@router.patch("/{category_id}")
def update_category(category_id: int, payload: CategoryUpdate, db: Session = Depends(get_db), _user=Depends(get_current_user)):
    category = db.get(Category, category_id)
    if category is None or category.deleted_at is not None:
        raise HTTPException(status_code=404, detail="Category not found")
    data = payload.model_dump(exclude_unset=True)
    if "slug" in data and data["slug"]:
        exists = db.query(Category).filter(Category.slug == data["slug"], Category.id != category_id, Category.deleted_at.is_(None)).first()
        if exists:
            raise HTTPException(status_code=400, detail="Slug already exists")
    if "parent_id" in data and data["parent_id"] is not None:
        if data["parent_id"] == category_id:
            raise HTTPException(status_code=400, detail="Cannot set category as its own parent")
        parent = db.get(Category, data["parent_id"])
        if parent is None or parent.deleted_at is not None:
            raise HTTPException(status_code=400, detail="Parent category not found")
    for key, value in data.items():
        setattr(category, key, value)
    db.commit()
    db.refresh(category)
    return ok(CategoryResponse.model_validate(category).model_dump(mode="json"))


@router.delete("/{category_id}")
def delete_category(category_id: int, db: Session = Depends(get_db), _user=Depends(get_current_user)):
    category = db.get(Category, category_id)
    if category is None or category.deleted_at is not None:
        raise HTTPException(status_code=404, detail="Category not found")
    category.deleted_at = datetime.now(UTC)
    db.add(category)
    db.commit()
    return ok({"status": "deleted"})
