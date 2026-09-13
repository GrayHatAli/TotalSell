"""Shared helper for reading uploaded Excel files.

Encapsulates the openpyxl read_only quirk: worksheets stream cells lazily
from the workbook's underlying ZIP archive, so every row must be
materialized BEFORE wb.close(). All import endpoints go through this
helper so that bug class cannot reappear in one of them.
"""
from io import BytesIO

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

ALLOWED_EXTS = (".xlsx", ".xlsm")


def read_spreadsheet_rows(file: UploadFile) -> list[tuple]:
    """Read all data rows (including the header) from an uploaded Excel file.

    Raises HTTPException(400) for a wrong file type, an unreadable file,
    a workbook without an active sheet, or an empty file.
    """
    filename = file.filename or ""
    if not filename.lower().endswith(ALLOWED_EXTS):
        raise HTTPException(status_code=400, detail="Only .xlsx and .xlsm files are supported")
    try:
        wb = load_workbook(BytesIO(file.file.read()), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read the Excel file")
    ws = wb.active
    if ws is None:
        wb.close()
        raise HTTPException(status_code=400, detail="Excel workbook has no active sheet")
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise HTTPException(status_code=400, detail="Excel file is empty")
    return rows